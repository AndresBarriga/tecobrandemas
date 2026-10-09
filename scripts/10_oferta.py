"""Anuncios recientes por distrito y barrio (€/m²) a partir de la serie del Ayuntamiento de Madrid.

Serie 4.3.21.D «Evolución del precio de oferta de alquiler de la vivienda (€/m²) por Distrito, Barrio y
Mes» (Banco de Datos del Ayuntamiento; elaboración del Ayuntamiento a partir de datos de Idealista).
Es la única fuente de oferta del producto: no se usa Idealista directamente.

Proceso mensual, a mano: se descarga el Excel (o su CSV) y se deja en data/raw/ con el mes del último dato en
el nombre (oferta_AAAA-MM.xlsx u oferta_AAAA-MM.csv). Sin argumentos se usa el más reciente.

Formato de la serie: año, distrito, barrio y una columna por mes («Mayo», «Junio»…); decimales con coma; «..»
= sin dato; una fila «Total Ciudad de Madrid»; filas de distrito («14. Moratalaz», distrito = barrio) y de
barrio («146. Vinateros»); al pie, la fuente y las equivalencias de barrios de Idealista con los del BOAM
(ya vienen aplicadas en los códigos, no hay nada que mapear).

Reglas (docs/decisiones.md):
  · mes = el último con dato en alguno de los distritos;
  · distrito: su valor del último mes con dato (cada valor lleva su mes);
  · barrio: solo si tiene dato en el último mes y en el anterior; se usa el valor del último. Si no, el
    producto usa el distrito (nunca otro barrio);
  · un valor que no se puede leer detiene el script: nunca se adivina.

Salida: data/processed/oferta_madrid.json
Uso: .venv/bin/python scripts/10_oferta.py [data/raw/oferta_AAAA-MM.xlsx|.csv]
"""

import csv
import json
import re
import sys
from datetime import date
from pathlib import Path

from comun import PROCESSED, RAW, escribir_json

SERIE = "4.3.21.D"
FUENTE = (
    "Ayuntamiento de Madrid, Banco de Datos, serie 4.3.21.D "
    "(elaboración del Ayuntamiento a partir de datos de Idealista)"
)
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
# €/m² al mes: fuera de este rango es un error de lectura, no un dato
MIN_M2, MAX_M2 = 5.0, 60.0


def celdas(ruta: Path) -> list[list[str]]:
    """Todas las filas de la hoja como texto, de un .csv o de un .xlsx (openpyxl)."""
    if ruta.suffix.lower() == ".csv":
        with open(ruta, encoding="utf-8-sig", newline="") as f:
            return [[c.strip() for c in fila] for fila in csv.reader(f)]
    from openpyxl import load_workbook

    hoja = load_workbook(ruta, read_only=True, data_only=True).worksheets[0]
    return [["" if c is None else str(c).strip() for c in fila] for fila in hoja.iter_rows(values_only=True)]


def valor(texto: str, donde: str) -> float | None:
    if texto in ("", ".."):
        return None
    limpio = texto.replace(",", ".")
    if not re.fullmatch(r"\d+(\.\d+)?", limpio):
        raise SystemExit(f"Valor ilegible «{texto}» en {donde}")
    v = float(limpio)
    if not MIN_M2 <= v <= MAX_M2:
        raise SystemExit(f"Valor fuera de rango ({v} €/m²) en {donde}")
    return v


def mes_del_nombre(ruta: Path) -> str | None:
    m = re.search(r"(\d{4})-(\d{2})", ruta.name)
    return f"{m.group(1)}-{m.group(2)}" if m else None


def elegir_entrada() -> Path:
    if len(sys.argv) > 1:
        return Path(sys.argv[1])
    candidatas = sorted(RAW.glob("oferta_*-*.csv")) + sorted(RAW.glob("oferta_*-*.xlsx"))
    if not candidatas:
        raise SystemExit(f"No hay ningún oferta_AAAA-MM.csv|xlsx en {RAW}")
    return max(candidatas, key=lambda p: mes_del_nombre(p) or "")


def main():
    entrada = elegir_entrada()
    filas = celdas(entrada)

    # Cabecera: la fila que trae nombres de mes
    i_cab = next((i for i, f in enumerate(filas) if sum(c.lower() in MESES for c in f) >= 1), None)
    if i_cab is None:
        raise SystemExit("No encuentro la fila de meses (Mayo, Junio…)")
    columnas = {j: MESES.index(c.lower()) + 1 for j, c in enumerate(filas[i_cab]) if c.lower() in MESES}

    ciudad: dict[int, float | None] = {}
    distritos: dict[str, dict[int, float | None]] = {}
    barrios: dict[str, dict[int, float | None]] = {}
    anios = set()
    for f in filas[i_cab + 1 :]:
        if len(f) < 3 or not re.fullmatch(r"\d{4}", f[0]):
            continue
        anios.add(int(f[0]))
        donde = f"{f[1]} / {f[2]}"
        por_mes = {mes: valor(f[j], f"{donde}, {MESES[mes - 1]}") for j, mes in columnas.items() if j < len(f)}
        if f[1].startswith("Total"):
            ciudad = por_mes
            continue
        # «14. Moratalaz» → 14; «171. Villaverde Alto, Casco Histórico de Villaverde» → 171
        codigo = f[2].split(".", 1)[0].strip()
        if not codigo.isdigit():
            raise SystemExit(f"Código ilegible en «{f[2]}»")
        if len(codigo) == 2:
            distritos[codigo] = por_mes
        elif len(codigo) == 3:
            barrios[codigo] = por_mes
        else:
            raise SystemExit(f"Código inesperado «{codigo}»")
    if len(anios) != 1:
        raise SystemExit(f"Se esperaba un solo año y hay {sorted(anios)}")
    anio = anios.pop()

    # Cobertura: los 131 barrios y 21 distritos oficiales (la misma tabla que usa el producto)
    oficiales = json.loads((PROCESSED / "seccion_barrio.json").read_text(encoding="utf-8"))["barrios"]
    if set(barrios) != set(oficiales):
        raise SystemExit(f"Barrios distintos de los oficiales: faltan {sorted(set(oficiales) - set(barrios))}, sobran {sorted(set(barrios) - set(oficiales))}")
    if len(distritos) != 21:
        raise SystemExit(f"Se esperaban 21 distritos y hay {len(distritos)}")

    # Meses con dato en algún distrito (el último y el anterior son los que cuentan)
    con_dato = sorted(m for m in columnas.values() if any(d.get(m) is not None for d in distritos.values()))
    if not con_dato:
        raise SystemExit("La serie no trae ningún dato")
    ultimo = con_dato[-1]
    anterior = con_dato[-2] if len(con_dato) > 1 else None
    mes = f"{anio}-{ultimo:02d}"
    esperado = mes_del_nombre(entrada)
    if esperado and esperado != mes:
        raise SystemExit(f"El nombre dice {esperado} pero el último mes con dato es {mes}")

    def clave(m: int) -> str:
        return f"{anio}-{m:02d}"

    salida_d = {}
    for cod, por_mes in sorted(distritos.items()):
        meses_ok = [m for m in con_dato if por_mes.get(m) is not None]
        if not meses_ok:
            continue  # sin dato en ningún mes: el producto no muestra la línea
        salida_d[cod] = {"v": por_mes[meses_ok[-1]], "mes": clave(meses_ok[-1])}

    salida_b = {}
    for cod, por_mes in sorted(barrios.items()):
        # Dos meses seguidos con dato: el último y el anterior
        if anterior is not None and por_mes.get(ultimo) is not None and por_mes.get(anterior) is not None:
            salida_b[cod] = {"v": por_mes[ultimo], "mes": mes}

    salida = {
        "serie": SERIE,
        "fuente": FUENTE,
        "mes": mes,
        "meses_con_dato": [clave(m) for m in con_dato],
        "ciudad": ciudad.get(ultimo),
        "distritos": salida_d,
        "barrios": salida_b,
        "fecha_extraccion": date.today().isoformat(),
    }

    print(f"{entrada.name}: mes {mes} (anterior {clave(anterior) if anterior else '—'})")
    print(f"  distritos con dato: {len(salida_d)} de {len(distritos)}")
    print(f"  barrios con dato en los dos últimos meses: {len(salida_b)} de {len(barrios)}")
    solo_ultimo = [c for c, m in barrios.items() if m.get(ultimo) is not None and c not in salida_b]
    sin_ultimo = [c for c, m in barrios.items() if anterior and m.get(anterior) is not None and m.get(ultimo) is None]
    print(f"  barrios con dato solo en {clave(ultimo)} (no se usan): {sorted(solo_ultimo)}")
    print(f"  barrios con dato en {clave(anterior) if anterior else '—'} y no en {clave(ultimo)}: {sorted(sin_ultimo)}")
    escribir_json(PROCESSED / "oferta_madrid.json", salida, compacto=False)


if __name__ == "__main__":
    main()
